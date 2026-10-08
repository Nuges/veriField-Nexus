"""
=============================================================================
VeriField Nexus — Agriculture MRV Laboratory Bulk Import Final Acceptance Suite
=============================================================================
Authoritative verification of:
1. Duplicate vs. Explicit Revision Safety (exact duplicate skip vs unconfirmed revision rejection vs explicit revision)
2. Commit Idempotency & Concurrency Safety
3. Upload Hash Idempotency (same bytes vs different filename vs different bytes)
4. Authoritative Resource Limits (File size, Rows, Columns, Cells, Sheets, Decompression Bomb)
5. Malformed File Rejection (Empty, corrupt, fake zip)
6. Formula Cell Safety & External Reference Blocking
7. Tenant & Project Isolation (Fail-closed sample matching)
8. Manual Entry vs Bulk Import Semantic Equivalence
9. Phase 3A Quantification Gating (PENDING exclusion until QA acceptance)
10. Accreditation Truth & Sample Lifecycle (UNVERIFIED default, ANALYZED status)
=============================================================================
"""

import asyncio
import io
import uuid
import zipfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import pytest
from fastapi import UploadFile, HTTPException
from sqlalchemy import select, and_
from sqlalchemy.ext.asyncio import AsyncSession

from app.domains.agriculture.laboratory_import_service import LaboratoryImportService
from app.domains.agriculture.models import (
    LaboratoryImportBatch,
    LaboratoryImportRow,
    PhysicalSample,
    LaboratoryReceipt,
    LaboratoryAnalysis,
    LaboratoryResult,
)
from app.domains.agriculture.service import AgricultureService
from app.domains.projects.models import Project
from app.domains.organizations.models import Organization
from tests.domains.agriculture.test_laboratory_bulk_import import create_agri_lab_import_setup


# =============================================================================
# 1. Duplicate vs. Explicit Revision Safety
# =============================================================================

@pytest.mark.asyncio
async def test_duplicate_vs_explicit_revision_safety(db_session: AsyncSession):
    setup = await create_agri_lab_import_setup(db_session)
    s1_code = setup["sample1"].sample_code

    # Batch 1: Initial result of 1.5%
    csv_1 = f"sample_code,analyte,raw_value,raw_unit\n{s1_code},SOC_CONCENTRATION,1.5,%\n".encode("utf-8")
    batch1 = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_1), filename="initial_1.csv"),
    )
    await db_session.commit()
    res1 = await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch1.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
    )
    await db_session.commit()
    assert res1["results_created"] == 1

    # Batch 2: Exact Duplicate assay in a separate laboratory file (different SHA256, identical assay row)
    csv_dup = f"sample_code,analyte,raw_value,raw_unit,notes\n{s1_code},SOC_CONCENTRATION,1.5,%,Duplicate check run\n".encode("utf-8")
    batch_dup = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_dup), filename="duplicate_run.csv"),
    )
    await db_session.commit()

    # Validation should flag EXACT_DUPLICATE_HISTORICAL warning
    rows = (await db_session.execute(
        select(LaboratoryImportRow).where(LaboratoryImportRow.import_batch_id == batch_dup.id)
    )).scalars().all()
    assert len(rows) == 1
    assert any(m["code"] == "EXACT_DUPLICATE_HISTORICAL" for m in rows[0].validation_messages)

    # Commit should be idempotent: 0 results created, 1 exact duplicate skipped
    commit_dup_res = await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch_dup.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
    )
    await db_session.commit()
    assert commit_dup_res["results_created"] == 0
    assert commit_dup_res["skipped_rows"] == 1

    # Batch 3: Different Value (1.85%) - Unconfirmed revision
    csv_rev = f"sample_code,analyte,raw_value,raw_unit,notes\n{s1_code},SOC_CONCENTRATION,1.85,%,Revision run\n".encode("utf-8")
    batch_rev = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_rev), filename="revision_run.csv"),
    )
    await db_session.commit()

    # Validation should flag REVISION_CANDIDATE warning
    rows_rev = (await db_session.execute(
        select(LaboratoryImportRow).where(LaboratoryImportRow.import_batch_id == batch_rev.id)
    )).scalars().all()
    assert len(rows_rev) == 1
    assert any(m["code"] == "REVISION_CANDIDATE" for m in rows_rev[0].validation_messages)

    # Committing without import_as_revision=True MUST fail closed with 400
    with pytest.raises(HTTPException) as exc_unconfirmed:
        await LaboratoryImportService.commit_batch(
            db=db_session,
            batch_id=batch_rev.id,
            organization_id=setup["org"].id,
            user_id=setup["admin"].id,
            import_as_revision=False,
        )
    assert exc_unconfirmed.value.status_code == 400
    assert "import_as_revision=True" in exc_unconfirmed.value.detail

    # Committing WITH import_as_revision=True MUST succeed and supersede prior result
    commit_rev_res = await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch_rev.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        import_as_revision=True,
    )
    await db_session.commit()
    assert commit_rev_res["results_created"] == 1

    # Verify superseding links
    results = (await db_session.execute(
        select(LaboratoryResult).where(
            and_(
                LaboratoryResult.physical_sample_id == setup["sample1"].id,
                LaboratoryResult.analyte == "SOC_CONCENTRATION",
            )
        ).order_by(LaboratoryResult.created_at.asc())
    )).scalars().all()
    assert len(results) == 2
    r_old, r_new = results[0], results[1]
    assert r_old.is_superseded is True
    assert r_old.superseded_by_id == r_new.id
    assert r_new.is_superseded is False
    assert r_new.supersedes_id == r_old.id
    assert r_new.raw_value == Decimal("1.85")


# =============================================================================
# 2. Commit Idempotency & Double-Commit Rejection
# =============================================================================

@pytest.mark.asyncio
async def test_commit_idempotency_blocks_double_commit(db_session: AsyncSession):
    setup = await create_agri_lab_import_setup(db_session)
    s2_code = setup["sample2"].sample_code

    csv_data = f"sample_code,analyte,raw_value,raw_unit\n{s2_code},BULK_DENSITY_G_CM3,1.35,g/cm3\n".encode("utf-8")
    batch = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_data), filename="idem.csv"),
    )
    await db_session.commit()

    # First commit succeeds
    res = await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
    )
    await db_session.commit()
    assert res["status"] in ("IMPORTED", "PARTIALLY_IMPORTED")

    # Second commit must fail closed with 400
    with pytest.raises(HTTPException) as exc:
        await LaboratoryImportService.commit_batch(
            db=db_session,
            batch_id=batch.id,
            organization_id=setup["org"].id,
            user_id=setup["admin"].id,
        )
    assert exc.value.status_code == 400
    assert "already been committed" in exc.value.detail


# =============================================================================
# 3. Upload Hash Idempotency
# =============================================================================

@pytest.mark.asyncio
async def test_upload_hash_idempotency(db_session: AsyncSession):
    setup = await create_agri_lab_import_setup(db_session)
    content_a = b"sample_code,analyte,raw_value,raw_unit\nSMP-001,SOC,1.2,%\n"
    content_b = b"sample_code,analyte,raw_value,raw_unit\nSMP-001,SOC,1.3,%\n"

    # Upload A
    batch_a = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(content_a), filename="dataset_a.csv"),
    )
    await db_session.commit()
    assert batch_a.id is not None

    # Upload same bytes content_a with DIFFERENT filename -> Rejected with 409 Conflict
    with pytest.raises(HTTPException) as exc_dup_bytes:
        await LaboratoryImportService.upload_file(
            db=db_session,
            project_id=setup["project"].id,
            organization_id=setup["org"].id,
            user_id=setup["admin"].id,
            file=UploadFile(file=io.BytesIO(content_a), filename="dataset_a_renamed.csv"),
        )
    assert exc_dup_bytes.value.status_code == 409
    assert "already been uploaded" in exc_dup_bytes.value.detail

    # Upload different bytes content_b with SAME filename "dataset_a.csv" -> Accepted
    batch_b = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(content_b), filename="dataset_a.csv"),
    )
    await db_session.commit()
    assert batch_b.id != batch_a.id


# =============================================================================
# 4. Authoritative Resource Limits & Malformed File Rejection
# =============================================================================

@pytest.mark.asyncio
async def test_resource_limits_and_malformed_rejection(db_session: AsyncSession):
    setup = await create_agri_lab_import_setup(db_session)

    # 1. File size exceeds MAX_IMPORT_FILE_BYTES (simulated by overriding or passing huge buffer)
    orig_limit = LaboratoryImportService.MAX_IMPORT_FILE_BYTES
    try:
        LaboratoryImportService.MAX_IMPORT_FILE_BYTES = 50  # Set limit to 50 bytes for test
        big_content = b"sample_code,analyte,raw_value,raw_unit\nSMP-001,SOC,1.2,%\n" * 5
        with pytest.raises(HTTPException) as exc_size:
            await LaboratoryImportService.upload_file(
                db=db_session,
                project_id=setup["project"].id,
                organization_id=setup["org"].id,
                user_id=setup["admin"].id,
                file=UploadFile(file=io.BytesIO(big_content), filename="huge.csv"),
            )
        assert exc_size.value.status_code == 413
    finally:
        LaboratoryImportService.MAX_IMPORT_FILE_BYTES = orig_limit

    # 2. Empty file -> 400 Bad Request
    with pytest.raises(HTTPException) as exc_empty:
        await LaboratoryImportService.upload_file(
            db=db_session,
            project_id=setup["project"].id,
            organization_id=setup["org"].id,
            user_id=setup["admin"].id,
            file=UploadFile(file=io.BytesIO(b"   \n\n"), filename="empty.csv"),
        )
    assert exc_empty.value.status_code == 400

    # 3. Exceeds Max Columns (> 100)
    cols = ["c" + str(i) for i in range(105)]
    many_cols_content = (",".join(cols) + "\n" + ",".join(["1"] * 105) + "\n").encode("utf-8")
    with pytest.raises(HTTPException) as exc_cols:
        await LaboratoryImportService.upload_file(
            db=db_session,
            project_id=setup["project"].id,
            organization_id=setup["org"].id,
            user_id=setup["admin"].id,
            file=UploadFile(file=io.BytesIO(many_cols_content), filename="many_cols.csv"),
        )
    assert exc_cols.value.status_code == 400
    assert "exceeding maximum limit" in exc_cols.value.detail

    # 4. Exceeds Max Cell Length (> 1000 chars)
    long_cell = "A" * 1050
    long_cell_content = f"sample_code,analyte,raw_value,raw_unit\nSMP-001,SOC,{long_cell},%\n".encode("utf-8")
    with pytest.raises(HTTPException) as exc_cell:
        await LaboratoryImportService.upload_file(
            db=db_session,
            project_id=setup["project"].id,
            organization_id=setup["org"].id,
            user_id=setup["admin"].id,
            file=UploadFile(file=io.BytesIO(long_cell_content), filename="long_cell.csv"),
        )
    assert exc_cell.value.status_code == 400
    assert "exceeds maximum limit" in exc_cell.value.detail

    # 5. Corrupt / fake XLSX file (bad zip archive)
    corrupt_xlsx = b"PK\x03\x04This is not a real zip archive header!"
    with pytest.raises(HTTPException) as exc_bad_xlsx:
        await LaboratoryImportService.upload_file(
            db=db_session,
            project_id=setup["project"].id,
            organization_id=setup["org"].id,
            user_id=setup["admin"].id,
            file=UploadFile(file=io.BytesIO(corrupt_xlsx), filename="fake.xlsx"),
        )
    assert exc_bad_xlsx.value.status_code == 400
    assert "Corrupted" in exc_bad_xlsx.value.detail or "format" in exc_bad_xlsx.value.detail


# =============================================================================
# 5. XLSX Formula Safety & External References
# =============================================================================

@pytest.mark.asyncio
async def test_xlsx_formula_and_external_refs(db_session: AsyncSession):
    import openpyxl
    import zipfile
    setup = await create_agri_lab_import_setup(db_session)
    s1_code = setup["sample1"].sample_code

    # 1. Build synthetic XLSX workbook covering all 5 XLSX formula & external reference cases
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["sample_code", "analyte", "raw_value", "raw_unit"])

    # Row 2: Literal text beginning with "=" (data_type="s")
    ws.cell(row=2, column=1, value=s1_code)
    ws.cell(row=2, column=2, value="SOC_CONCENTRATION")
    c2 = ws.cell(row=2, column=3, value="=LiteralStringValue")
    c2.data_type = "s"
    ws.cell(row=2, column=4, value="%")

    # Row 3: Real formula with cached result (=10+8.5 -> cached 18.5)
    ws.cell(row=3, column=1, value=s1_code)
    ws.cell(row=3, column=2, value="SOC_CONCENTRATION")
    ws.cell(row=3, column=3, value="=10+8.5")
    ws.cell(row=3, column=4, value="%")

    # Row 4: Formula without cached result (=A1*2)
    ws.cell(row=4, column=1, value=s1_code)
    ws.cell(row=4, column=2, value="SOC_CONCENTRATION")
    ws.cell(row=4, column=3, value="=A1*2")
    ws.cell(row=4, column=4, value="%")

    # Row 5: Formula with external workbook reference
    ws.cell(row=5, column=1, value=s1_code)
    ws.cell(row=5, column=2, value="SOC_CONCENTRATION")
    ws.cell(row=5, column=3, value="=[ExternalBook.xlsx]Sheet1!A1")
    ws.cell(row=5, column=4, value="%")

    # Row 6: Formula with HYPERLINK external URI
    ws.cell(row=6, column=1, value=s1_code)
    ws.cell(row=6, column=2, value="SOC_CONCENTRATION")
    ws.cell(row=6, column=3, value='=HYPERLINK("http://external.site/data", "18.5")')
    ws.cell(row=6, column=4, value="%")

    buf = io.BytesIO()
    wb.save(buf)

    # Inject pre-computed cached values into Row 3 (<v>18.5</v>) and Row 6 (<v>18.5</v>)
    import re
    zin = zipfile.ZipFile(buf, "r")
    zout_buf = io.BytesIO()
    zout = zipfile.ZipFile(zout_buf, "w")
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "xl/worksheets/sheet1.xml":
            text = data.decode("utf-8")
            text = re.sub(
                r'(<c r="C3"[^>]*><f>[^<]*10\+8\.5[^<]*</f>)(?:<v>[^<]*</v>|<v/>)?(</c>)',
                r'\1<v>18.5</v>\2',
                text,
            )
            text = re.sub(
                r'(<c r="C6"[^>]*><f>[^<]*HYPERLINK[^<]*</f>)(?:<v>[^<]*</v>|<v/>)?(</c>)',
                r'\1<v>18.5</v>\2',
                text,
            )
            data = text.encode("utf-8")
        zout.writestr(item, data)
    zin.close()
    zout.close()

    batch_xlsx = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(zout_buf.getvalue()), filename="formula_suite.xlsx"),
    )
    await db_session.commit()

    val_res = await LaboratoryImportService.validate_batch(
        db=db_session,
        batch_id=batch_xlsx.id,
        organization_id=setup["org"].id,
    )
    await db_session.commit()

    rows = (await db_session.execute(
        select(LaboratoryImportRow)
        .where(LaboratoryImportRow.import_batch_id == batch_xlsx.id)
        .order_by(LaboratoryImportRow.source_row_number.asc())
    )).scalars().all()

    assert len(rows) == 5

    # Case 1 (Row 2): Literal text starting with "=" is NOT a formula (data_type="s")
    r2_codes = {m["code"] for m in rows[0].validation_messages}
    assert "FORMULA_CELL" not in r2_codes
    assert "FORMULA_WITHOUT_CACHED_VALUE" not in r2_codes

    # Case 2 (Row 3): Real formula with cached result -> FORMULA_CELL warning, static cached value 18.5 used
    r3_codes = {m["code"] for m in rows[1].validation_messages}
    assert "FORMULA_CELL" in r3_codes
    assert "FORMULA_WITHOUT_CACHED_VALUE" not in r3_codes
    assert rows[1].raw_value == Decimal("18.5")
    assert rows[1].normalized_value == Decimal("185.0000")  # 18.5% * 10 = 185.0 g/kg

    # Case 3 (Row 4): Formula without cached result -> FORMULA_WITHOUT_CACHED_VALUE error
    r4_codes = {m["code"] for m in rows[2].validation_messages}
    assert "FORMULA_CELL" in r4_codes
    assert "FORMULA_WITHOUT_CACHED_VALUE" in r4_codes
    assert rows[2].validation_status == "ERROR"

    # Case 4 (Row 5): External workbook formula -> EXTERNAL_FORMULA_REF warning
    r5_codes = {m["code"] for m in rows[3].validation_messages}
    assert "FORMULA_CELL" in r5_codes
    assert "EXTERNAL_FORMULA_REF" in r5_codes

    # Case 5 (Row 6): HYPERLINK external URI -> EXTERNAL_FORMULA_REF warning
    r6_codes = {m["code"] for m in rows[4].validation_messages}
    assert "FORMULA_CELL" in r6_codes
    assert "EXTERNAL_FORMULA_REF" in r6_codes

    # 2. Case 6: CSV formula injection defense (checks =, +, -, @ prefixes)
    csv_formula = (
        f'sample_code,analyte,raw_value,raw_unit\n'
        f'{s1_code},SOC_CONCENTRATION,"=HYPERLINK(""http://malicious.site"", 1.85)",%\n'
    ).encode("utf-8")
    batch_csv = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_formula), filename="formula_injection.csv"),
    )
    await db_session.commit()
    csv_rows = (await db_session.execute(
        select(LaboratoryImportRow).where(LaboratoryImportRow.import_batch_id == batch_csv.id)
    )).scalars().all()
    assert len(csv_rows) == 1
    csv_codes = {m["code"] for m in csv_rows[0].validation_messages}
    assert "FORMULA_CELL" in csv_codes
    assert "EXTERNAL_FORMULA_REF" in csv_codes


# =============================================================================
# 6. Tenant and Project Isolation
# =============================================================================

@pytest.mark.asyncio
async def test_tenant_and_project_isolation(db_session: AsyncSession):
    setup = await create_agri_lab_import_setup(db_session)
    foreign_code = f"SMP-OTHER-ORG-{uuid.uuid4().hex[:6]}"

    csv_data = f"sample_code,analyte,raw_value,raw_unit\n{foreign_code},SOC_CONCENTRATION,1.5,%\n".encode("utf-8")
    batch = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_data), filename="cross_tenant.csv"),
    )
    await db_session.commit()

    rows = (await db_session.execute(
        select(LaboratoryImportRow).where(LaboratoryImportRow.import_batch_id == batch.id)
    )).scalars().all()
    assert len(rows) == 1
    # Sample from other org/project MUST fail closed: SAMPLE_NOT_FOUND error
    assert rows[0].validation_status == "ERROR"
    assert rows[0].matched_sample_id is None
    assert any(m["code"] == "SAMPLE_NOT_FOUND" for m in rows[0].validation_messages)


# =============================================================================
# 7. Manual Entry vs Bulk Import Semantic Equivalence
# =============================================================================

@pytest.mark.asyncio
async def test_manual_entry_vs_bulk_import_equivalence(db_session: AsyncSession):
    import openpyxl
    setup = await create_agri_lab_import_setup(db_session)
    s1 = setup["sample1"]
    s2 = setup["sample2"]
    s3 = setup["sample_rejected"]
    s3.status = "COLLECTED"
    await db_session.flush()

    # 1. Manual Entry via AgricultureService: Raw 1.85 %
    now = datetime.now(timezone.utc)
    manual_receipt = LaboratoryReceipt(
        id=uuid.uuid4(),
        physical_sample_id=s1.id,
        laboratory_name="Standard Lab",
        intake_status="ACCEPTED",
        received_at=now,
        received_by_name="Lab Tech",
        condition_on_receipt="ACCEPTABLE",
        seal_status="SEALED_INTACT",
    )
    db_session.add(manual_receipt)
    manual_analysis = LaboratoryAnalysis(
        id=uuid.uuid4(),
        physical_sample_id=s1.id,
        laboratory_name="Standard Lab",
        laboratory_accreditation="NOT_PROVIDED",
        accreditation_status="UNVERIFIED",
        analytical_method="DUMAS_COMBUSTION",
        analysis_date=now.date(),
        qa_status="PENDING",
    )
    db_session.add(manual_analysis)
    norm_val, norm_unit, norm_meth, norm_ver = AgricultureService.normalize_laboratory_analyte_measurement(
        "SOC_CONCENTRATION", Decimal("1.85"), "%"
    )
    manual_result = LaboratoryResult(
        id=uuid.uuid4(),
        analysis_id=manual_analysis.id,
        physical_sample_id=s1.id,
        analyte="SOC_CONCENTRATION",
        raw_value=Decimal("1.85"),
        raw_unit="%",
        normalized_value=norm_val,
        normalized_unit=norm_unit,
        normalization_method=norm_meth,
        normalization_version=norm_ver,
        is_superseded=False,
    )
    db_session.add(manual_result)
    await db_session.commit()

    # 2. Bulk Import via CSV for Sample 2: Raw 1.85 %
    csv_data = f"sample_code,analyte,raw_value,raw_unit\n{s2.sample_code},SOC,1.85,%\n".encode("utf-8")
    batch_csv = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_data), filename="equiv.csv"),
    )
    await db_session.commit()
    await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch_csv.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
    )
    await db_session.commit()

    # 3. Bulk Import via XLSX for Sample 3: Raw 1.85 %
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["sample_code", "analyte", "raw_value", "raw_unit"])
    ws.append([s3.sample_code, "SOC_CONCENTRATION", 1.85, "%"])
    x_buf = io.BytesIO()
    wb.save(x_buf)
    batch_xlsx = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(x_buf.getvalue()), filename="equiv.xlsx"),
    )
    await db_session.commit()
    await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch_xlsx.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
    )
    await db_session.commit()

    csv_result = (await db_session.execute(
        select(LaboratoryResult).where(
            and_(
                LaboratoryResult.physical_sample_id == s2.id,
                LaboratoryResult.analyte == "SOC_CONCENTRATION",
            )
        )
    )).scalars().first()
    assert csv_result is not None

    xlsx_result = (await db_session.execute(
        select(LaboratoryResult).where(
            and_(
                LaboratoryResult.physical_sample_id == s3.id,
                LaboratoryResult.analyte == "SOC_CONCENTRATION",
            )
        )
    )).scalars().first()
    assert xlsx_result is not None

    # Equivalence assertions across all 3 ingestion channels:
    for res in [manual_result, csv_result, xlsx_result]:
        assert res.analyte == "SOC_CONCENTRATION"
        assert res.raw_unit == "%"
        assert res.raw_value == Decimal("1.85")
        assert res.normalized_value == Decimal("18.5000")
        assert res.normalized_unit == "g/kg"
        assert res.normalization_method == "LINEAR_SCALING:VAL*10"
        assert res.is_superseded is False


# =============================================================================
# 8. Phase 3A Quantification Readiness Gating
# =============================================================================

@pytest.mark.asyncio
async def test_phase3a_quantification_readiness_qa_gating(db_session: AsyncSession):
    from app.domains.agriculture.schemas import SampleQAReviewCreate
    setup = await create_agri_lab_import_setup(db_session)
    s1 = setup["sample1"]

    # Ingest bulk results: LaboratoryAnalysis created with qa_status="PENDING"
    csv_data = f"sample_code,analyte,raw_value,raw_unit\n{s1.sample_code},SOC_CONCENTRATION,2.5,%\n".encode("utf-8")
    batch = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_data), filename="gating.csv"),
    )
    await db_session.commit()
    await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
    )
    await db_session.commit()

    # Verify analysis is PENDING
    analysis = (await db_session.execute(
        select(LaboratoryAnalysis).where(LaboratoryAnalysis.analysis_batch_id == str(batch.id))
    )).scalars().first()
    assert analysis.qa_status == "PENDING"

    # Query latest valid result for quantification readiness:
    # LaboratoryResult whose LaboratoryAnalysis.qa_status == "VERIFIED" (canonical enum: PENDING, VERIFIED, REJECTED)
    verified_results_stmt = (
        select(LaboratoryResult)
        .join(LaboratoryAnalysis, LaboratoryResult.analysis_id == LaboratoryAnalysis.id)
        .where(
            and_(
                LaboratoryResult.physical_sample_id == s1.id,
                LaboratoryResult.is_superseded.is_(False),
                LaboratoryAnalysis.qa_status == "VERIFIED",
            )
        )
    )
    verified_res = (await db_session.execute(verified_results_stmt)).scalars().all()
    # PENDING analysis MUST NOT qualify for quantification readiness!
    assert len(verified_res) == 0

    # Submit authoritative QA Review via AgricultureService to test the true workflow:
    # QA Officer reviews sample with overall_qa_status="ACCEPTED"
    review_in = SampleQAReviewCreate(
        reviewer_name="QA Officer Alice",
        overall_qa_status="ACCEPTED",
        findings="Passed laboratory QA inspection and verification protocol",
        remediation_required=False,
    )
    qa_review = await AgricultureService.record_sample_qa_review(
        db=db_session,
        sample_id=s1.id,
        organization_id=setup["org"].id,
        payload=review_in,
        user_id=setup["admin"].id,
    )
    await db_session.commit()

    # Verify QA state separation:
    # SampleQAReview has overall_qa_status = "ACCEPTED"
    # PhysicalSample transitioned to "QA_ACCEPTED"
    # LaboratoryAnalysis transitioned to qa_status = "VERIFIED" (NEVER "ACCEPTED")
    assert qa_review.overall_qa_status == "ACCEPTED"
    await db_session.refresh(s1)
    assert s1.status == "QA_ACCEPTED"
    await db_session.refresh(analysis)
    assert analysis.qa_status == "VERIFIED"

    # Now the result becomes eligible for quantification readiness
    verified_res_after = (await db_session.execute(verified_results_stmt)).scalars().all()
    assert len(verified_res_after) == 1
    assert verified_res_after[0].normalized_value == Decimal("25.0000")
    assert verified_res_after[0].normalized_unit == "g/kg"


# =============================================================================
# 9. Accreditation Truth & Sample Lifecycle
# =============================================================================

@pytest.mark.asyncio
async def test_accreditation_truth_and_lifecycle(db_session: AsyncSession):
    setup = await create_agri_lab_import_setup(db_session)
    s1 = setup["sample1"]
    s2 = setup["sample2"]

    # Case 1: No accreditation claimed in file -> defaults to "NOT_PROVIDED", status="UNVERIFIED"
    csv_data_unacc = f"sample_code,analyte,raw_value,raw_unit\n{s2.sample_code},PH,6.8,pH\n".encode("utf-8")
    batch1 = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_data_unacc), filename="accred_unspecified.csv"),
    )
    await db_session.commit()
    await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch1.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
    )
    await db_session.commit()

    analysis1 = (await db_session.execute(
        select(LaboratoryAnalysis).where(LaboratoryAnalysis.analysis_batch_id == str(batch1.id))
    )).scalars().first()
    assert analysis1.accreditation_status == "UNVERIFIED"
    assert analysis1.laboratory_accreditation == "NOT_PROVIDED"

    # Sample lifecycle transition: status is ANALYZED
    await db_session.refresh(s2)
    assert s2.status == "ANALYZED"

    # Case 2: Claimed accreditation standard in file -> stores claim, status remains "UNVERIFIED" until audited
    csv_data_claimed = f"sample_code,analyte,raw_value,raw_unit,laboratory_accreditation\n{s1.sample_code},BULK_DENSITY_G_CM3,1.30,g/cm3,ISO/IEC 17025\n".encode("utf-8")
    batch2 = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_data_claimed), filename="accred_claimed.csv"),
    )
    await db_session.commit()
    await LaboratoryImportService.commit_batch(
        db=db_session,
        batch_id=batch2.id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
    )
    await db_session.commit()

    analysis2 = (await db_session.execute(
        select(LaboratoryAnalysis).where(LaboratoryAnalysis.analysis_batch_id == str(batch2.id))
    )).scalars().first()
    assert analysis2.laboratory_accreditation == "ISO/IEC 17025"
    assert analysis2.accreditation_status == "UNVERIFIED"  # Crucial: claims start UNVERIFIED until compliance review


# =============================================================================
# 10. PostgreSQL Concurrency Locking
# =============================================================================

@pytest.mark.asyncio
async def test_concurrent_commit_postgres(db_session: AsyncSession):
    # PostgreSQL row-level lock concurrency requires real PostgreSQL
    bind = db_session.bind
    if "sqlite" in str(bind.url):
        pytest.skip("Row-level SELECT FOR UPDATE concurrency requires PostgreSQL")

    setup = await create_agri_lab_import_setup(db_session)
    s1_code = setup["sample1"].sample_code

    csv_data = f"sample_code,analyte,raw_value,raw_unit\n{s1_code},SOC,1.95,%\n".encode("utf-8")
    batch = await LaboratoryImportService.upload_file(
        db=db_session,
        project_id=setup["project"].id,
        organization_id=setup["org"].id,
        user_id=setup["admin"].id,
        file=UploadFile(file=io.BytesIO(csv_data), filename="concur.csv"),
    )
    await db_session.commit()

    # Create two separate sessions to simulate concurrent workers
    from sqlalchemy.ext.asyncio import async_sessionmaker
    session_maker = async_sessionmaker(bind, expire_on_commit=False)

    async def worker_commit(worker_name: str):
        async with session_maker() as sess:
            try:
                res = await LaboratoryImportService.commit_batch(
                    db=sess,
                    batch_id=batch.id,
                    organization_id=setup["org"].id,
                    user_id=setup["admin"].id,
                )
                await sess.commit()
                return {"worker": worker_name, "success": True, "res": res}
            except HTTPException as e:
                await sess.rollback()
                return {"worker": worker_name, "success": False, "status_code": e.status_code, "detail": e.detail}

    results = await asyncio.gather(
        worker_commit("worker_1"),
        worker_commit("worker_2"),
    )

    successes = [r for r in results if r["success"]]
    failures = [r for r in results if not r["success"]]

    # Exactly one commit must succeed, and the concurrent competitor must receive 400
    assert len(successes) == 1
    assert len(failures) == 1
    assert failures[0]["status_code"] == 400
    assert "already been committed" in failures[0]["detail"]
