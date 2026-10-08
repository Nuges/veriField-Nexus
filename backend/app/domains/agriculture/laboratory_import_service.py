"""
=============================================================================
VeriField Nexus — Agriculture MRV Laboratory Bulk Data Import Service
=============================================================================
Provides production-grade ingestion, parsing, staging, automated column
detection, strict scientific validation, sample matching, deduplication,
formula injection defense, and transactional commit of bulk laboratory
datasets (.csv, .xlsx).

Integrates natively with:
- PhysicalSample & SamplingCampaign
- Canonical Unit Normalization (AgricultureService)
- Immutable Evidence Vault (Evidence table & SHA-256 integrity)
- QA Review & Laboratory Analysis Lifecycle
=============================================================================
"""

import csv
import io
import os
import re
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select, and_, or_
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from fastapi import HTTPException, UploadFile, status

from app.core.object_storage import S3StorageManager
from app.domains.documents.security import DocumentSecurityValidator
from app.domains.evidence.models import Evidence
from app.domains.agriculture.models import (
    LaboratoryImportBatch,
    LaboratoryImportRow,
    PhysicalSample,
    LaboratoryReceipt,
    LaboratoryAnalysis,
    LaboratoryResult,
    SamplingCampaign,
)
from app.domains.agriculture.service import AgricultureService
from app.domains.agriculture.schemas import ColumnMappingConfig


# Canonical analyte alias definitions
ANALYTE_ALIASES: Dict[str, str] = {
    # Soil Organic Carbon
    "SOC": "SOC_CONCENTRATION",
    "SOC_CONCENTRATION": "SOC_CONCENTRATION",
    "SOC_STOCK_PCT": "SOC_CONCENTRATION",
    "SOC_PCT": "SOC_CONCENTRATION",
    "TOTAL_ORGANIC_CARBON": "SOC_CONCENTRATION",
    "TOTAL_ORGANIC_CARBON_G_KG": "SOC_CONCENTRATION",
    "SOIL_ORGANIC_CARBON": "SOC_CONCENTRATION",
    "ORGANIC_CARBON": "SOC_CONCENTRATION",
    "C_ORG": "SOC_CONCENTRATION",
    "TOC": "SOC_CONCENTRATION",
    # Bulk Density
    "BULK_DENSITY": "BULK_DENSITY_G_CM3",
    "BULK_DENSITY_G_CM3": "BULK_DENSITY_G_CM3",
    "CORE_BULK_DENSITY": "BULK_DENSITY_G_CM3",
    "BD": "BULK_DENSITY_G_CM3",
    "DRY_BULK_DENSITY": "BULK_DENSITY_G_CM3",
    # Texture & Fractions
    "COARSE_FRAGMENTS": "COARSE_FRAGMENTS_PCT",
    "COARSE_FRAGMENTS_PCT": "COARSE_FRAGMENTS_PCT",
    "GRAVEL": "COARSE_FRAGMENTS_PCT",
    "GRAVEL_PCT": "COARSE_FRAGMENTS_PCT",
    "STONES": "COARSE_FRAGMENTS_PCT",
    "SAND": "SAND_PCT",
    "SAND_PCT": "SAND_PCT",
    "SILT": "SILT_PCT",
    "SILT_PCT": "SILT_PCT",
    "CLAY": "CLAY_PCT",
    "CLAY_PCT": "CLAY_PCT",
    "MOISTURE": "MOISTURE_PCT",
    "MOISTURE_PCT": "MOISTURE_PCT",
    # Chemical / Physical
    "PH": "PH",
    "SOIL_PH": "PH",
    "PH_H2O": "PH",
    "PH_CACL2": "PH",
    "EC": "ELECTRICAL_CONDUCTIVITY_DS_M",
    "ELECTRICAL_CONDUCTIVITY": "ELECTRICAL_CONDUCTIVITY_DS_M",
    "ELECTRICAL_CONDUCTIVITY_DS_M": "ELECTRICAL_CONDUCTIVITY_DS_M",
    "SALINITY": "ELECTRICAL_CONDUCTIVITY_DS_M",
    "TOTAL_NITROGEN": "TOTAL_NITROGEN_PCT",
    "TOTAL_NITROGEN_PCT": "TOTAL_NITROGEN_PCT",
    "TN": "TOTAL_NITROGEN_PCT",
    "N_TOT": "TOTAL_NITROGEN_PCT",
}

COLUMN_NAME_CANDIDATES: Dict[str, List[str]] = {
    "sample_code": ["sample_code", "sample_id", "sample_identifier", "sampleid", "field_sample_id", "sample", "code"],
    "analyte": ["analyte", "parameter", "property", "test_parameter", "measurement", "variable", "compound"],
    "value": ["raw_value", "value", "result", "reading", "concentration", "val", "measurement_value"],
    "unit": ["raw_unit", "unit", "units", "uom", "dimension"],
    "depth_from": ["depth_from_cm", "depth_from", "top_depth", "min_depth", "depth_upper", "start_depth", "depth_top"],
    "depth_to": ["depth_to_cm", "depth_to", "bottom_depth", "max_depth", "depth_lower", "end_depth", "depth_bottom"],
    "analysis_date": ["analysis_date", "date_analyzed", "test_date", "date", "run_date"],
    "method": ["method", "analytical_method", "test_method", "protocol", "standard_method", "technique"],
    "lab_sample_id": ["lab_sample_id", "lab_id", "laboratory_sample_id", "lab_barcode", "internal_id"],
    "notes": ["notes", "remarks", "comments", "observation", "qa_notes"],
}


class LaboratoryImportService:
    """Service handling bulk laboratory data ingestion, validation, preview, and transactional commit."""

    # Authoritative Backend Resource Limits & Security Constraints
    MAX_IMPORT_FILE_BYTES: int = 50 * 1024 * 1024  # 50 MB
    MAX_IMPORT_ROWS: int = 5000
    MAX_IMPORT_COLUMNS: int = 100
    MAX_IMPORT_CELL_LENGTH: int = 1000
    MAX_IMPORT_SHEETS: int = 10
    MAX_UNCOMPRESSED_ARCHIVE_BYTES: int = 200 * 1024 * 1024  # 200 MB safe decompression limit

    @staticmethod
    def sanitize_formula_injection(val: Any) -> str:
        """
        Escapes spreadsheet formula execution characters (=, +, -, @) to neutralize
        CSV / Excel formula injection vulnerabilities.
        """
        if val is None:
            return ""
        s = str(val).strip()
        if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
            return f"'{s}"
        return s

    @classmethod
    def generate_template(cls, file_format: str = "csv") -> Tuple[bytes, str, str]:
        """
        Generates a standard Laboratory Ingestion Template with sample realistic soil rows.
        Returns: (file_bytes, filename, media_type)
        """
        headers = [
            "sample_code",
            "analyte",
            "raw_value",
            "raw_unit",
            "depth_from_cm",
            "depth_to_cm",
            "analysis_date",
            "method",
            "lab_sample_id",
            "notes",
        ]
        sample_rows = [
            [
                "SMP-SOIL-001",
                "SOC_CONCENTRATION",
                1.85,
                "%",
                0,
                30,
                "2026-09-15",
                "DUMAS_COMBUSTION",
                "LAB-2026-0901",
                "Dry combustion elemental analyzer (ISO 10694)",
            ],
            [
                "SMP-SOIL-001",
                "BULK_DENSITY_G_CM3",
                1.32,
                "g/cm3",
                0,
                30,
                "2026-09-15",
                "CORE_METHOD",
                "LAB-2026-0902",
                "100cm3 ring core oven-dried at 105C (ISO 11272)",
            ],
            [
                "SMP-SOIL-001",
                "PH",
                6.45,
                "pH",
                0,
                30,
                "2026-09-15",
                "1:5_H2O",
                "LAB-2026-0903",
                "Electrometric in 1:5 soil:water suspension (ISO 10390)",
            ],
            [
                "SMP-SOIL-002",
                "SOC_CONCENTRATION",
                24.50,
                "g/kg",
                0,
                30,
                "2026-09-15",
                "DUMAS_COMBUSTION",
                "LAB-2026-0904",
                "Analytical duplicate run (ISO 10694)",
            ],
            [
                "SMP-SOIL-002",
                "BULK_DENSITY_G_CM3",
                1.28,
                "g/cm3",
                0,
                30,
                "2026-09-15",
                "CORE_METHOD",
                "LAB-2026-0905",
                "Intact soil core sample (ISO 11272)",
            ],
            [
                "SMP-SOIL-002",
                "COARSE_FRAGMENTS_PCT",
                4.2,
                "%",
                0,
                30,
                "2026-09-15",
                "GRAVIMETRIC_SIEVING",
                "LAB-2026-0906",
                "Sieved through 2mm round hole sieve",
            ],
        ]

        if file_format.lower() == "xlsx":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Lab_Results_Import"

            # Header styling
            header_fill = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='DDDDDD'),
                right=Side(style='thin', color='DDDDDD'),
                top=Side(style='thin', color='DDDDDD'),
                bottom=Side(style='thin', color='DDDDDD')
            )

            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for row_idx, r_data in enumerate(sample_rows, start=2):
                for col_idx, val in enumerate(r_data, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.border = thin_border
                    if isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="left")

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

            out_stream = io.BytesIO()
            wb.save(out_stream)
            content = out_stream.getvalue()
            return (
                content,
                "verifield_laboratory_import_template.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        # CSV format
        out_stream = io.StringIO()
        writer = csv.writer(out_stream, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(headers)
        for r in sample_rows:
            writer.writerow(r)
        content = out_stream.getvalue().encode("utf-8")
        return (content, "verifield_laboratory_import_template.csv", "text/csv")

    @classmethod
    def suggest_column_mappings(cls, headers: List[str]) -> Dict[str, Any]:
        """
        Inspects spreadsheet header names and suggests best-fit column mappings
        based on standard agronomic MRV conventions.
        """
        normalized_map = {}
        for h in headers:
            clean = re.sub(r'[^a-zA-Z0-9]', '_', h).strip('_').lower()
            normalized_map[clean] = h

        detected = {
            "sample_code_column": "",
            "analyte_column": "",
            "value_column": "",
            "unit_column": "",
            "depth_from_column": "",
            "depth_to_column": "",
            "analysis_date_column": "",
            "method_column": "",
            "lab_sample_id_column": "",
            "notes_column": "",
        }

        for target, candidates in COLUMN_NAME_CANDIDATES.items():
            field_key = f"{target}_column" if not target.endswith("_column") else target
            # Try exact match first
            for cand in candidates:
                if cand in normalized_map:
                    detected[field_key] = normalized_map[cand]
                    break
            # Try fuzzy substring if still unset
            if not detected[field_key]:
                for norm_h, orig_h in normalized_map.items():
                    for cand in candidates:
                        if cand in norm_h:
                            detected[field_key] = orig_h
                            break
                    if detected[field_key]:
                        break

        # Defaults if candidates weren't found
        if not detected["sample_code_column"] and "sample_code" in headers:
            detected["sample_code_column"] = "sample_code"
        if not detected["value_column"] and "raw_value" in headers:
            detected["value_column"] = "raw_value"
        if not detected["analyte_column"] and "analyte" in headers:
            detected["analyte_column"] = "analyte"

        return detected

    @classmethod
    async def upload_file(
        cls,
        db: AsyncSession,
        project_id: uuid.UUID,
        organization_id: uuid.UUID,
        user_id: uuid.UUID,
        file: UploadFile,
        laboratory_name: Optional[str] = None,
        sampling_campaign_id: Optional[uuid.UUID] = None,
    ) -> LaboratoryImportBatch:
        """
        Securely ingests, validates, hashes, stages rows, and stores evidence for a bulk lab file.
        """
        # Validate security, MIME, magic bytes, sanitize filename, compute SHA-256
        content, clean_filename, mime_type, sha256_hash = await DocumentSecurityValidator.validate_and_read(file)
        ext = os.path.splitext(clean_filename)[1].lower()
        file_type = "CSV" if ext == ".csv" else ("XLSX" if ext == ".xlsx" else "OTHER")

        # Check authoritative file size limit
        if len(content) > cls.MAX_IMPORT_FILE_BYTES:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"File exceeds maximum allowed size of {cls.MAX_IMPORT_FILE_BYTES // (1024 * 1024)}MB.",
            )

        # Check for duplicate file upload in same project by cryptographic content hash (SHA-256)
        dup_stmt = select(LaboratoryImportBatch).where(
            and_(
                LaboratoryImportBatch.project_id == project_id,
                LaboratoryImportBatch.file_sha256 == sha256_hash,
            )
        )
        existing_batch = (await db.execute(dup_stmt)).scalars().first()
        if existing_batch:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"This exact laboratory file has already been uploaded to this project (Batch ID: {existing_batch.id}).",
            )

        # Store file in Object Storage or static uploads directory
        storage_mgr = S3StorageManager()
        storage_res = await storage_mgr.upload_file(content, clean_filename, mime_type)
        file_uri = storage_res.get("signed_url") or f"local://uploads/{sha256_hash}{ext}"

        # Create immutable Evidence record
        evidence = Evidence(
            id=uuid.uuid4(),
            activity_id=project_id,
            file_uri=file_uri,
            file_hash=sha256_hash,
            evidence_type="LABORATORY_REPORT",
            metadata_json={
                "original_filename": clean_filename,
                "file_size_bytes": len(content),
                "mime_type": mime_type,
                "laboratory_name": laboratory_name,
                "sampling_campaign_id": str(sampling_campaign_id) if sampling_campaign_id else None,
            },
            status="PENDING",
            uploaded_by=user_id,
        )
        db.add(evidence)
        await db.flush()

        # Parse rows from file
        parsed_rows: List[Tuple[str, int, Dict[str, Any]]] = []  # (sheet_name, row_num, raw_payload)
        headers_found: List[str] = []

        if file_type == "CSV":
            try:
                # Detect encoding and delimiter
                text_content = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text_content = content.decode("latin-1")
                except Exception as ue:
                    raise HTTPException(status_code=400, detail=f"Invalid file encoding: {str(ue)}")

            if not text_content.strip():
                raise HTTPException(status_code=400, detail="CSV file is empty.")

            stream = io.StringIO(text_content)
            sample_chunk = text_content[:2048]
            delimiter = ","
            try:
                sniffer = csv.Sniffer()
                dialect = sniffer.sniff(sample_chunk)
                delimiter = dialect.delimiter
            except Exception:
                delimiter = ","

            stream.seek(0)
            reader = csv.reader(stream, delimiter=delimiter)
            try:
                all_lines = list(reader)
            except Exception as ce:
                raise HTTPException(status_code=400, detail=f"Malformed CSV formatting: {str(ce)}")

            if not all_lines or not all_lines[0]:
                raise HTTPException(status_code=400, detail="CSV file contains no data or headers.")

            headers_found = [h.strip() for h in all_lines[0]]
            if len(headers_found) > cls.MAX_IMPORT_COLUMNS:
                raise HTTPException(
                    status_code=400,
                    detail=f"CSV file contains {len(headers_found)} columns, exceeding maximum limit of {cls.MAX_IMPORT_COLUMNS}.",
                )

            for idx, r in enumerate(all_lines[1:], start=2):
                if not any(bool(c.strip()) for c in r):
                    continue  # skip completely blank lines

                if len(parsed_rows) >= cls.MAX_IMPORT_ROWS:
                    raise HTTPException(
                        status_code=400,
                        detail=f"File exceeds maximum allowed rows limit of {cls.MAX_IMPORT_ROWS}.",
                    )

                row_dict = {}
                for col_idx, col_name in enumerate(headers_found):
                    raw_val = r[col_idx].strip() if col_idx < len(r) else ""
                    if len(raw_val) > cls.MAX_IMPORT_CELL_LENGTH:
                        raise HTTPException(
                            status_code=400,
                            detail=f"Cell content at row {idx}, column '{col_name}' exceeds maximum limit of {cls.MAX_IMPORT_CELL_LENGTH} characters.",
                        )
                    row_dict[col_name] = raw_val
                parsed_rows.append(("Sheet1", idx, row_dict))

        elif file_type == "XLSX":
            import zipfile
            external_relationships: List[Dict[str, str]] = []
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as z:
                    total_uncompressed = sum(info.file_size for info in z.infolist())
                    if total_uncompressed > cls.MAX_UNCOMPRESSED_ARCHIVE_BYTES:
                        raise HTTPException(
                            status_code=400,
                            detail="Workbook uncompressed size exceeds maximum safety limit (potential decompression bomb).",
                        )
                    total_compressed = sum(info.compress_size for info in z.infolist())
                    if total_compressed > 0 and (total_uncompressed / total_compressed > 100) and total_uncompressed > 10 * 1024 * 1024:
                        raise HTTPException(
                            status_code=400,
                            detail="Abnormal archive compression ratio detected (potential decompression bomb).",
                        )
                    # Inspect OOXML package relationships for external targets and links
                    for name in z.namelist():
                        name_lower = name.lower()
                        if any(k in name_lower for k in ["externallink", "externalbook", "ddelink", "olelink"]):
                            external_relationships.append({
                                "file": name,
                                "type": "EXTERNAL_BOOK_LINK",
                                "detail": "Workbook contains external workbook or DDE/OLE link part.",
                            })
                        elif name.endswith(".rels"):
                            rel_content = z.read(name).decode("utf-8", errors="ignore")
                            if 'TargetMode="External"' in rel_content or "relationships/externalLink" in rel_content:
                                external_relationships.append({
                                    "file": name,
                                    "type": "EXTERNAL_RELATIONSHIP",
                                    "detail": "Relationship with TargetMode=External detected in OOXML package.",
                                })
            except zipfile.BadZipFile:
                raise HTTPException(status_code=400, detail="Corrupted or invalid XLSX archive format.")
            except HTTPException:
                raise
            except Exception as ze:
                raise HTTPException(status_code=400, detail=f"Failed to inspect XLSX archive: {str(ze)}")

            # PASS A: Safe non-executing formula inspection (data_only=False, read_only=True)
            try:
                wb_formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=True)
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to parse Excel workbook: {str(e)}")

            if len(wb_formulas.sheetnames) > cls.MAX_IMPORT_SHEETS:
                raise HTTPException(
                    status_code=400,
                    detail=f"Workbook contains {len(wb_formulas.sheetnames)} sheets, exceeding maximum limit of {cls.MAX_IMPORT_SHEETS}.",
                )

            sheet_name = wb_formulas.sheetnames[0]
            ws_formulas = wb_formulas[sheet_name]
            try:
                raw_formula_rows = list(ws_formulas.iter_rows())
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to inspect formula cells: {str(e)}")
            wb_formulas.close()

            # PASS B: Static cached value retrieval (data_only=True, read_only=True)
            try:
                wb_values = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to parse cached values in Excel workbook: {str(e)}")

            ws_values = wb_values[sheet_name]
            try:
                raw_value_rows = list(ws_values.iter_rows(values_only=True))
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to read sheet data: {str(e)}")
            wb_values.close()

            if not raw_value_rows or len(raw_value_rows) < 2:
                raise HTTPException(status_code=400, detail="Excel spreadsheet contains no data rows.")

            headers_found = [str(c or "").strip() for c in raw_value_rows[0] if c is not None]
            if len(headers_found) > cls.MAX_IMPORT_COLUMNS:
                raise HTTPException(
                    status_code=400,
                    detail=f"Spreadsheet contains {len(headers_found)} columns, exceeding maximum limit of {cls.MAX_IMPORT_COLUMNS}.",
                )

            for idx, r in enumerate(raw_value_rows[1:], start=2):
                if not any(c is not None and str(c).strip() for c in r):
                    continue  # skip empty lines

                if len(parsed_rows) >= cls.MAX_IMPORT_ROWS:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Spreadsheet exceeds maximum allowed rows limit of {cls.MAX_IMPORT_ROWS}.",
                    )

                formula_cells_row = raw_formula_rows[idx - 1] if (idx - 1) < len(raw_formula_rows) else []

                row_dict: Dict[str, Any] = {}
                formula_meta: Dict[str, Any] = {}

                for col_idx, col_name in enumerate(headers_found):
                    val = r[col_idx] if col_idx < len(r) else None
                    f_cell = formula_cells_row[col_idx] if col_idx < len(formula_cells_row) else None

                    is_formula = False
                    formula_expr = None
                    has_external_ref = False
                    has_cached_val = val is not None and str(val).strip() != ""

                    if f_cell is not None:
                        # Identify real formula cells strictly by openpyxl cell data_type == 'f'
                        # or formula expression starting with '=' when data_type is not explicitly string literal 's'
                        f_val_str = str(f_cell.value or "").strip()
                        if f_cell.data_type == 'f' or (f_cell.data_type != 's' and f_val_str.startswith('=')):
                            is_formula = True
                            formula_expr = f_val_str
                            expr_upper = formula_expr.upper()
                            if any(marker in expr_upper for marker in ["[", "]", "HYPERLINK", "DDE", "EXEC", "CMD", "WEBSERVICE", "SHELL", ".XLS"]):
                                has_external_ref = True

                    if is_formula or (f_cell is not None and f_cell.data_type == 's'):
                        formula_meta[col_name] = {
                            "is_formula": is_formula,
                            "formula_expr": formula_expr,
                            "has_cached_val": has_cached_val,
                            "has_external_ref": has_external_ref,
                            "data_type": getattr(f_cell, "data_type", None) if f_cell is not None else None,
                        }

                    if isinstance(val, (datetime, date)):
                        val_str = val.isoformat()
                    elif val is not None:
                        val_str = str(val).strip()
                    elif is_formula and not has_cached_val:
                        val_str = formula_expr or ""
                    else:
                        val_str = ""

                    if len(val_str) > cls.MAX_IMPORT_CELL_LENGTH:
                        raise HTTPException(
                            status_code=400,
                            detail=f"Cell content at row {idx}, column '{col_name}' exceeds maximum limit of {cls.MAX_IMPORT_CELL_LENGTH} characters.",
                        )
                    row_dict[col_name] = val_str

                if formula_meta:
                    row_dict["_formula_meta"] = formula_meta
                if external_relationships:
                    row_dict["_package_external_refs"] = external_relationships

                parsed_rows.append((sheet_name, idx, row_dict))

        if not parsed_rows:
            raise HTTPException(status_code=400, detail="No readable data rows found in uploaded file.")

        suggested_mapping = cls.suggest_column_mappings(headers_found)

        # Create LaboratoryImportBatch
        batch = LaboratoryImportBatch(
            id=uuid.uuid4(),
            organization_id=organization_id,
            project_id=project_id,
            sampling_campaign_id=sampling_campaign_id,
            laboratory_name=laboratory_name,
            original_filename=clean_filename,
            file_type=file_type,
            file_size_bytes=len(content),
            file_sha256=sha256_hash,
            evidence_id=evidence.id,
            status="UPLOADED",
            source_type="FILE_IMPORT",
            uploaded_by=user_id,
            mapping_version="V1.0",
            mapping_config=suggested_mapping,
            total_rows=len(parsed_rows),
            valid_rows=0,
            warning_rows=0,
            error_rows=0,
            imported_rows=0,
            skipped_rows=0,
            error_summary=[],
        )
        db.add(batch)
        await db.flush()

        # Stage raw rows
        for sheet, row_num, payload in parsed_rows:
            staged_row = LaboratoryImportRow(
                id=uuid.uuid4(),
                import_batch_id=batch.id,
                organization_id=organization_id,
                project_id=project_id,
                source_sheet_name=sheet,
                source_row_number=row_num,
                raw_row_payload=payload,
                mapped_payload={},
                validation_status="VALID",
                validation_messages=[],
            )
            db.add(staged_row)

        await db.flush()

        # Auto-run initial validation using suggested mapping
        await cls.validate_batch(db, batch.id, organization_id, suggested_mapping, laboratory_name, sampling_campaign_id)
        await db.refresh(batch)
        return batch

    @classmethod
    async def validate_batch(
        cls,
        db: AsyncSession,
        batch_id: uuid.UUID,
        organization_id: uuid.UUID,
        mapping_config: Optional[Dict[str, Any]] = None,
        laboratory_name: Optional[str] = None,
        sampling_campaign_id: Optional[uuid.UUID] = None,
    ) -> Tuple[LaboratoryImportBatch, List[LaboratoryImportRow]]:
        """
        Executes strict scientific and relational validation over all staged rows in a batch.
        Updates row status, validation messages, matched samples, and canonical conversions.
        """
        # Load batch
        batch_stmt = select(LaboratoryImportBatch).where(
            and_(
                LaboratoryImportBatch.id == batch_id,
                LaboratoryImportBatch.organization_id == organization_id,
            )
        )
        batch = (await db.execute(batch_stmt)).scalars().first()
        if not batch:
            raise HTTPException(status_code=404, detail="Laboratory import batch not found.")

        if batch.status in ("IMPORTED", "PARTIALLY_IMPORTED"):
            raise HTTPException(
                status_code=400,
                detail=f"Cannot re-validate batch in status '{batch.status}'. Batch is already committed.",
            )

        if mapping_config:
            batch.mapping_config = mapping_config
        if laboratory_name is not None:
            batch.laboratory_name = laboratory_name
        if sampling_campaign_id is not None:
            batch.sampling_campaign_id = sampling_campaign_id

        cfg = batch.mapping_config or {}
        sample_col = cfg.get("sample_code_column", "sample_code")
        analyte_col = cfg.get("analyte_column", "analyte")
        analyte_const = cfg.get("analyte_constant")
        val_col = cfg.get("value_column", "raw_value")
        unit_col = cfg.get("unit_column", "raw_unit")
        unit_const = cfg.get("unit_constant")
        depth_from_col = cfg.get("depth_from_column", "depth_from_cm")
        depth_to_col = cfg.get("depth_to_column", "depth_to_cm")
        date_col = cfg.get("analysis_date_column", "analysis_date")
        method_col = cfg.get("method_column", "method")
        lab_id_col = cfg.get("lab_sample_id_column", "lab_sample_id")
        notes_col = cfg.get("notes_column", "notes")

        # Fetch all staged rows for this batch
        rows_stmt = (
            select(LaboratoryImportRow)
            .where(LaboratoryImportRow.import_batch_id == batch.id)
            .order_by(LaboratoryImportRow.source_row_number.asc())
        )
        rows = (await db.execute(rows_stmt)).scalars().all()

        # Pre-fetch all PhysicalSamples for this project to eliminate N+1 queries
        samples_stmt = (
            select(PhysicalSample)
            .where(PhysicalSample.project_id == batch.project_id)
            .options(selectinload(PhysicalSample.sampling_point))
        )
        samples_list = (await db.execute(samples_stmt)).scalars().all()
        sample_by_code: Dict[str, PhysicalSample] = {
            s.sample_code.strip().upper(): s for s in samples_list if s.sample_code
        }

        # Pre-fetch existing active LaboratoryResults for samples in this project
        sample_ids = [s.id for s in samples_list]
        existing_active_results_map: Dict[Tuple[uuid.UUID, str], LaboratoryResult] = {}
        if sample_ids:
            results_stmt = select(LaboratoryResult).where(
                and_(
                    LaboratoryResult.physical_sample_id.in_(sample_ids),
                    LaboratoryResult.is_superseded.is_(False),
                )
            )
            for res_obj in (await db.execute(results_stmt)).scalars().all():
                if res_obj.analyte:
                    existing_active_results_map[(res_obj.physical_sample_id, res_obj.analyte.upper())] = res_obj

        valid_count = 0
        warning_count = 0
        error_count = 0
        all_error_summaries: List[Dict[str, Any]] = []
        batch_seen_keys: Set[Tuple[str, str, str, str]] = set()

        for row in rows:
            payload = row.raw_row_payload or {}
            messages: List[Dict[str, Any]] = []
            mapped: Dict[str, Any] = {}

            # 1. Sample Code Extraction & Relational Matching
            raw_sample_code = str(payload.get(sample_col, "")).strip()
            mapped["sample_code"] = raw_sample_code
            matched_sample: Optional[PhysicalSample] = None

            if not raw_sample_code:
                messages.append({
                    "severity": "ERROR",
                    "code": "MISSING_SAMPLE_CODE",
                    "message": f"Column '{sample_col}' is empty or missing.",
                })
            else:
                lookup_key = raw_sample_code.upper()
                if lookup_key not in sample_by_code:
                    messages.append({
                        "severity": "ERROR",
                        "code": "SAMPLE_NOT_FOUND",
                        "message": f"Physical sample '{raw_sample_code}' does not exist in project.",
                    })
                else:
                    matched_sample = sample_by_code[lookup_key]
                    row.matched_sample_id = matched_sample.id
                    row.matched_sample_code = matched_sample.sample_code

                    # Lifecycle check
                    if matched_sample.status in ("REJECTED_BY_LAB", "CANCELLED", "DISCARDED"):
                        messages.append({
                            "severity": "ERROR",
                            "code": "SAMPLE_REJECTED",
                            "message": f"Physical sample '{raw_sample_code}' has status '{matched_sample.status}'. Results cannot be recorded for rejected samples.",
                        })

            # 2. Depth Interval Cross-Check
            raw_df = payload.get(depth_from_col, "")
            raw_dt = payload.get(depth_to_col, "")
            df_val: Optional[float] = None
            dt_val: Optional[float] = None

            if str(raw_df).strip():
                try:
                    df_val = float(str(raw_df).strip())
                    mapped["depth_from_cm"] = df_val
                except ValueError:
                    messages.append({
                        "severity": "ERROR",
                        "code": "INVALID_DEPTH",
                        "message": f"Invalid upper depth value '{raw_df}'. Must be numeric.",
                    })

            if str(raw_dt).strip():
                try:
                    dt_val = float(str(raw_dt).strip())
                    mapped["depth_to_cm"] = dt_val
                except ValueError:
                    messages.append({
                        "severity": "ERROR",
                        "code": "INVALID_DEPTH",
                        "message": f"Invalid lower depth value '{raw_dt}'. Must be numeric.",
                    })

            if df_val is not None and dt_val is not None:
                if dt_val <= df_val:
                    messages.append({
                        "severity": "ERROR",
                        "code": "INVALID_DEPTH_INTERVAL",
                        "message": f"Lower depth ({dt_val} cm) must be greater than upper depth ({df_val} cm).",
                    })

            # Compare against physical sample depth via associated sampling point
            if matched_sample and df_val is not None and dt_val is not None:
                pt = matched_sample.sampling_point
                s_df = float(pt.depth_from_cm) if pt and pt.depth_from_cm is not None else None
                s_dt = float(pt.depth_to_cm) if pt and pt.depth_to_cm is not None else None
                if s_df is not None and s_dt is not None:
                    if abs(df_val - s_df) > 0.05 or abs(dt_val - s_dt) > 0.05:
                        messages.append({
                            "severity": "WARNING",
                            "code": "DEPTH_MISMATCH",
                            "message": f"Spreadsheet depth ({df_val}-{dt_val} cm) differs from planned sample depth ({s_df}-{s_dt} cm).",
                        })

            # 3. Analyte Resolution & Normalization
            raw_analyte = str(payload.get(analyte_col, "")).strip() or (analyte_const or "").strip()
            mapped["analyte"] = raw_analyte
            canonical_analyte: Optional[str] = None

            if not raw_analyte:
                messages.append({
                    "severity": "ERROR",
                    "code": "MISSING_ANALYTE",
                    "message": "Analyte is missing. Specify an analyte column or default constant.",
                })
            else:
                an_clean = re.sub(r'[^a-zA-Z0-9]', '_', raw_analyte).strip('_').upper()
                canonical_analyte = ANALYTE_ALIASES.get(an_clean, an_clean)
                row.canonical_analyte = canonical_analyte
                mapped["canonical_analyte"] = canonical_analyte

            # 4. Raw Value Parsing
            raw_val_str = str(payload.get(val_col, "")).strip()
            mapped["raw_value_str"] = raw_val_str
            dec_value: Optional[Decimal] = None

            if not raw_val_str:
                messages.append({
                    "severity": "ERROR",
                    "code": "MISSING_VALUE",
                    "message": f"Measurement value in column '{val_col}' is missing.",
                })
            else:
                formula_meta = payload.get("_formula_meta", {}).get(val_col, {})
                package_external_refs = payload.get("_package_external_refs", [])

                if batch.file_type == "XLSX" or formula_meta:
                    if formula_meta.get("is_formula"):
                        f_expr = formula_meta.get("formula_expr", "")
                        messages.append({
                            "severity": "WARNING",
                            "code": "FORMULA_CELL",
                            "message": f"XLSX formula '{f_expr}' detected. Dynamic formula execution is disabled; static cached value is used.",
                        })
                        if not formula_meta.get("has_cached_val"):
                            messages.append({
                                "severity": "ERROR",
                                "code": "FORMULA_WITHOUT_CACHED_VALUE",
                                "message": f"Formula '{f_expr}' contains no pre-evaluated cached result in workbook. Uncached formulas cannot be imported.",
                            })
                        if formula_meta.get("has_external_ref") or package_external_refs:
                            messages.append({
                                "severity": "WARNING",
                                "code": "EXTERNAL_FORMULA_REF",
                                "message": "Formula or workbook contains external references/links. External resolution is blocked (zero network access policy).",
                            })
                    elif batch.file_type == "XLSX" and raw_val_str.startswith("=") and formula_meta.get("data_type") != "s":
                        messages.append({
                            "severity": "WARNING",
                            "code": "FORMULA_CELL",
                            "message": f"XLSX formula '{raw_val_str}' detected. Dynamic formula execution is disabled.",
                        })
                        if any(marker in raw_val_str.upper() for marker in ["[", "]", "HYPERLINK", "DDE", "EXEC", "CMD", "WEBSERVICE", "SHELL", ".XLS"]):
                            messages.append({
                                "severity": "WARNING",
                                "code": "EXTERNAL_FORMULA_REF",
                                "message": "Formula contains external references/links. External resolution is blocked.",
                            })
                else:
                    # CSV formula injection defense (checks leading special characters =, +, -, @, \t, \r)
                    if raw_val_str and raw_val_str[0] in ("=", "+", "-", "@", "\t", "\r"):
                        messages.append({
                            "severity": "WARNING",
                            "code": "FORMULA_CELL",
                            "message": "Formula prefix detected in CSV cell. Dynamic formula execution is disabled.",
                        })
                        if any(bad in raw_val_str.upper() for bad in ["HTTP", "HYPERLINK", "DDE", "EXEC", "SHELL", ".XLS", "CMD"]):
                            messages.append({
                                "severity": "WARNING",
                                "code": "EXTERNAL_FORMULA_REF",
                                "message": "Formula contains external or dynamic references. External references are blocked.",
                            })

                if package_external_refs and not any(m["code"] == "EXTERNAL_FORMULA_REF" for m in messages):
                    messages.append({
                        "severity": "WARNING",
                        "code": "EXTERNAL_FORMULA_REF",
                        "message": "Workbook contains external package relationships. External resolution is blocked (zero network access policy).",
                    })

                # If formula has no cached result, do not attempt numeric conversion
                is_uncached_formula = (formula_meta.get("is_formula") and not formula_meta.get("has_cached_val")) or (
                    batch.file_type == "XLSX" and raw_val_str.startswith("=") and formula_meta.get("data_type") != "s" and not formula_meta.get("has_cached_val")
                )
                if is_uncached_formula:
                    dec_value = None
                else:
                    try:
                        dec_value = Decimal(raw_val_str)
                        row.raw_value = dec_value
                        mapped["raw_value"] = float(dec_value)
                    except InvalidOperation:
                        messages.append({
                            "severity": "ERROR",
                            "code": "NON_NUMERIC_VALUE",
                            "message": f"Measurement value '{raw_val_str}' is not a valid decimal number.",
                        })

            # 5. Raw Unit & Scientific Conversion
            raw_unit = str(payload.get(unit_col, "")).strip() or (unit_const or "").strip()
            row.raw_unit = raw_unit
            mapped["raw_unit"] = raw_unit

            if not raw_unit:
                messages.append({
                    "severity": "ERROR",
                    "code": "MISSING_UNIT",
                    "message": "Unit is missing. Specify a unit column or default constant.",
                })
            elif dec_value is not None and canonical_analyte:
                norm_val, norm_unit, norm_meth, norm_ver = AgricultureService.normalize_laboratory_analyte_measurement(
                    canonical_analyte, dec_value, raw_unit
                )

                if norm_meth == "UNCONVERTIBLE" or norm_val is None:
                    messages.append({
                        "severity": "ERROR",
                        "code": "UNCONVERTIBLE_UNIT",
                        "message": f"Cannot normalize unit '{raw_unit}' for analyte '{canonical_analyte}'. Supported units for SOC include %, g/kg, mg/kg.",
                    })
                else:
                    row.normalized_value = norm_val
                    row.normalized_unit = norm_unit
                    mapped["normalized_value"] = float(norm_val)
                    mapped["normalized_unit"] = norm_unit
                    mapped["normalization_method"] = norm_meth
                    mapped["normalization_version"] = norm_ver

                    # Plausibility & Domain Range Verification
                    val_float = float(norm_val)
                    if canonical_analyte == "SOC_CONCENTRATION":
                        if val_float < 0:
                            messages.append({
                                "severity": "ERROR",
                                "code": "NEGATIVE_SOC",
                                "message": f"SOC concentration cannot be negative ({val_float} g/kg).",
                            })
                        elif val_float > 500:
                            messages.append({
                                "severity": "WARNING",
                                "code": "EXTREME_SOC",
                                "message": f"High SOC concentration ({val_float} g/kg). Standard mineral soils are < 100 g/kg; peat/organic soils up to 500 g/kg.",
                            })
                    elif canonical_analyte == "BULK_DENSITY_G_CM3":
                        if val_float <= 0.1 or val_float > 2.65:
                            messages.append({
                                "severity": "ERROR",
                                "code": "OUT_OF_BOUNDS_BULK_DENSITY",
                                "message": f"Bulk density ({val_float} g/cm³) is outside physical soil limits (0.1 - 2.65 g/cm³).",
                            })
                    elif canonical_analyte == "PH":
                        if val_float < 2.0 or val_float > 12.0:
                            messages.append({
                                "severity": "ERROR",
                                "code": "OUT_OF_BOUNDS_PH",
                                "message": f"Soil pH ({val_float}) is outside valid scientific limits (2.0 - 12.0).",
                            })
                    elif canonical_analyte == "COARSE_FRAGMENTS_PCT":
                        if val_float < 0.0 or val_float > 100.0:
                            messages.append({
                                "severity": "ERROR",
                                "code": "OUT_OF_BOUNDS_COARSE_FRAGMENTS",
                                "message": f"Coarse fragments ({val_float}%) must be between 0% and 100%.",
                            })

            # 6. Duplicate Detection (Intra-batch & Historical)
            if raw_sample_code and canonical_analyte:
                dup_key = (
                    raw_sample_code.upper(),
                    canonical_analyte.upper(),
                    str(df_val or ""),
                    str(dt_val or ""),
                )
                if dup_key in batch_seen_keys:
                    messages.append({
                        "severity": "ERROR",
                        "code": "DUPLICATE_IN_BATCH",
                        "message": f"Duplicate assay for sample '{raw_sample_code}', analyte '{canonical_analyte}' in this batch.",
                    })
                else:
                    batch_seen_keys.add(dup_key)

                # Historical duplicate & revision candidate check
                if matched_sample and (matched_sample.id, canonical_analyte.upper()) in existing_active_results_map:
                    hist_res = existing_active_results_map[(matched_sample.id, canonical_analyte.upper())]
                    # Check if exact duplicate value
                    is_exact = False
                    if row.normalized_value is not None and hist_res.normalized_value is not None:
                        if abs(float(hist_res.normalized_value) - float(row.normalized_value)) < 1e-5 and str(hist_res.normalized_unit or "").upper() == str(row.normalized_unit or "").upper():
                            is_exact = True
                    elif dec_value is not None and hist_res.raw_value is not None:
                        if abs(float(hist_res.raw_value) - float(dec_value)) < 1e-5 and str(hist_res.raw_unit or "").upper() == str(raw_unit or "").upper():
                            is_exact = True

                    val_display = str(row.raw_value if row.raw_value is not None else raw_val_str)
                    if is_exact:
                        messages.append({
                            "severity": "WARNING",
                            "code": "EXACT_DUPLICATE_HISTORICAL",
                            "message": f"Sample '{raw_sample_code}' already has an active '{canonical_analyte}' result with identical value ({hist_res.raw_value} {hist_res.raw_unit}). Row will be treated idempotently and skipped on commit.",
                        })
                    else:
                        messages.append({
                            "severity": "WARNING",
                            "code": "REVISION_CANDIDATE",
                            "message": f"Sample '{raw_sample_code}' already has an active '{canonical_analyte}' result with different value ({hist_res.raw_value} {hist_res.raw_unit} vs {val_display} {raw_unit}). Must explicitly specify import_as_revision=True to commit this revision.",
                        })

            # Metadata pass-through
            mapped["analysis_date"] = payload.get(date_col, "")
            mapped["method"] = payload.get(method_col, "")
            mapped["lab_sample_id"] = payload.get(lab_id_col, "")
            mapped["notes"] = payload.get(notes_col, "")

            raw_accred = (
                str(payload.get(cfg.get("accreditation_column", "accreditation"), "")).strip()
                or str(payload.get("accreditation", "")).strip()
                or str(payload.get("laboratory_accreditation", "")).strip()
                or str(payload.get("accreditation_standard", "")).strip()
            )
            if raw_accred:
                mapped["laboratory_accreditation"] = raw_accred

            # Set status for row
            errors = [m for m in messages if m["severity"] == "ERROR"]
            warnings = [m for m in messages if m["severity"] == "WARNING"]

            if errors:
                row.validation_status = "ERROR"
                error_count += 1
            elif warnings:
                row.validation_status = "WARNING"
                warning_count += 1
            else:
                row.validation_status = "VALID"
                valid_count += 1

            row.validation_messages = messages
            row.mapped_payload = mapped

            for err in errors:
                all_error_summaries.append({
                    "row": row.source_row_number,
                    "sample_code": raw_sample_code,
                    "code": err["code"],
                    "message": err["message"],
                })

        batch.valid_rows = valid_count
        batch.warning_rows = warning_count
        batch.error_rows = error_count
        batch.status = "VALIDATED" if error_count == 0 else "VALIDATED_WITH_ERRORS"
        batch.error_summary = all_error_summaries[:50]  # Cap top 50 in summary JSON

        await db.flush()
        return batch, rows

    @classmethod
    async def commit_batch(
        cls,
        db: AsyncSession,
        batch_id: uuid.UUID,
        organization_id: uuid.UUID,
        user_id: uuid.UUID,
        import_valid_only: bool = False,
        laboratory_name: Optional[str] = None,
        notes: Optional[str] = None,
        import_as_revision: bool = False,
        revision_reason: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Transactionally commits validated rows into canonical LaboratoryAnalysis and
        LaboratoryResult domain records. Links evidence provenance and updates PhysicalSample status.
        Enforces strict duplicate idempotency (exact duplicates skipped) and explicit revision safety
        (different values require import_as_revision=True).
        """
        # Concurrency safety: acquire row lock
        batch_stmt = (
            select(LaboratoryImportBatch)
            .where(
                and_(
                    LaboratoryImportBatch.id == batch_id,
                    LaboratoryImportBatch.organization_id == organization_id,
                )
            )
            .with_for_update()
        )
        batch = (await db.execute(batch_stmt)).scalars().first()
        if not batch:
            raise HTTPException(status_code=404, detail="Laboratory import batch not found.")

        if batch.status in ("IMPORTED", "PARTIALLY_IMPORTED"):
            raise HTTPException(
                status_code=400,
                detail="Batch has already been committed to the canonical laboratory ledger.",
            )

        if batch.error_rows > 0 and not import_valid_only:
            raise HTTPException(
                status_code=400,
                detail=f"Batch has {batch.error_rows} error rows. Resolve errors or select 'import_valid_only=True' to proceed.",
            )

        # Fetch rows
        rows_stmt = (
            select(LaboratoryImportRow)
            .where(LaboratoryImportRow.import_batch_id == batch.id)
            .order_by(LaboratoryImportRow.source_row_number.asc())
        )
        rows = (await db.execute(rows_stmt)).scalars().all()

        eligible_rows = [r for r in rows if r.validation_status in ("VALID", "WARNING") and r.matched_sample_id]
        skipped_rows = [r for r in rows if r not in eligible_rows]

        if not eligible_rows:
            raise HTTPException(
                status_code=400,
                detail="No eligible valid rows available to import.",
            )

        lab_name = laboratory_name or batch.laboratory_name or "External Soil Laboratory"

        # Group eligible rows by matched sample
        sample_ids_set = {r.matched_sample_id for r in eligible_rows}
        samples_stmt = select(PhysicalSample).where(PhysicalSample.id.in_(sample_ids_set))
        samples_map = {s.id: s for s in (await db.execute(samples_stmt)).scalars().all()}

        # Verify or auto-create LaboratoryReceipt for matched samples so intake contract is satisfied
        receipts_stmt = select(LaboratoryReceipt).where(LaboratoryReceipt.physical_sample_id.in_(sample_ids_set))
        existing_receipts = {rc.physical_sample_id: rc for rc in (await db.execute(receipts_stmt)).scalars().all()}

        now_utc = datetime.now(timezone.utc)
        for s_id in sample_ids_set:
            if s_id not in existing_receipts:
                sample_obj = samples_map.get(s_id)
                receipt = LaboratoryReceipt(
                    id=uuid.uuid4(),
                    physical_sample_id=s_id,
                    laboratory_name=lab_name,
                    intake_status="ACCEPTED",
                    received_at=now_utc,
                    received_by_name=f"Bulk Importer (User {user_id})",
                    condition_on_receipt="ACCEPTABLE",
                    seal_status="SEALED_INTACT",
                    receipt_evidence_id=batch.evidence_id,
                )
                db.add(receipt)
                existing_receipts[s_id] = receipt

        await db.flush()

        # Group rows by sample
        rows_by_sample: Dict[uuid.UUID, List[LaboratoryImportRow]] = {}
        for r in eligible_rows:
            rows_by_sample.setdefault(r.matched_sample_id, []).append(r)

        analyses_created = 0
        results_created = 0
        exact_duplicates_skipped = 0
        samples_analyzed = set()

        for s_id, s_rows in rows_by_sample.items():
            sample_obj = samples_map.get(s_id)
            if not sample_obj:
                continue

            # Check rows for duplicate vs revision vs new
            rows_to_create_with_prev: List[Tuple[LaboratoryImportRow, Optional[LaboratoryResult]]] = []

            for r in s_rows:
                existing_active_stmt = select(LaboratoryResult).where(
                    and_(
                        LaboratoryResult.physical_sample_id == s_id,
                        LaboratoryResult.analyte == r.canonical_analyte,
                        LaboratoryResult.is_superseded.is_(False),
                    )
                )
                prev_res = (await db.execute(existing_active_stmt)).scalars().first()

                if prev_res:
                    is_exact = False
                    if r.normalized_value is not None and prev_res.normalized_value is not None:
                        if abs(float(prev_res.normalized_value) - float(r.normalized_value)) < 1e-5 and str(prev_res.normalized_unit or "").upper() == str(r.normalized_unit or "").upper():
                            is_exact = True
                    elif r.raw_value is not None and prev_res.raw_value is not None:
                        if abs(float(prev_res.raw_value) - float(r.raw_value)) < 1e-5 and str(prev_res.raw_unit or "").upper() == str(r.raw_unit or "").upper():
                            is_exact = True

                    if is_exact:
                        # Exact duplicate: skip creating redundant result (idempotent)
                        r.resulting_lab_result_id = prev_res.id
                        r.validation_status = "SKIPPED_DUPLICATE"
                        exact_duplicates_skipped += 1
                        continue

                    # Different value exists: check import_as_revision
                    if not import_as_revision:
                        raise HTTPException(
                            status_code=status.HTTP_400_BAD_REQUEST,
                            detail=(
                                f"Sample '{sample_obj.sample_code}' already has an active '{r.canonical_analyte}' result "
                                f"with different value ({prev_res.raw_value} {prev_res.raw_unit} vs {r.raw_value} {r.raw_unit}). "
                                "Overwriting requires explicit 'import_as_revision=True'."
                            ),
                        )

                rows_to_create_with_prev.append((r, prev_res))

            # Only create LaboratoryAnalysis if at least one result needs to be created
            if rows_to_create_with_prev:
                first_mapped = s_rows[0].mapped_payload or {}
                raw_date = first_mapped.get("analysis_date")
                parsed_date: Optional[date] = None
                if raw_date:
                    try:
                        if isinstance(raw_date, str):
                            parsed_date = date.fromisoformat(raw_date[:10])
                        elif isinstance(raw_date, date):
                            parsed_date = raw_date
                    except Exception:
                        parsed_date = now_utc.date()
                if not parsed_date:
                    parsed_date = now_utc.date()

                claimed_accred = first_mapped.get("laboratory_accreditation")
                raw_method = first_mapped.get("method") or first_mapped.get("analytical_method") or "DRY_COMBUSTION"
                method_str = str(raw_method).upper() if raw_method else "DRY_COMBUSTION"

                analysis = LaboratoryAnalysis(
                    id=uuid.uuid4(),
                    physical_sample_id=s_id,
                    laboratory_name=lab_name,
                    laboratory_accreditation=claimed_accred if claimed_accred else "NOT_PROVIDED",
                    accreditation_status="UNVERIFIED",  # Strictly preserve accreditation truth: claims start as UNVERIFIED until compliance verification
                    analysis_batch_id=str(batch.id),
                    analytical_method=method_str,
                    method_standard_code="ISO 10694 / ISO 11272",
                    analysis_date=parsed_date,
                    report_reference_number=batch.original_filename,
                    analyst_name=f"Bulk Importer (User {user_id})",
                    qa_status="PENDING",  # Strictly preserve SoD: imported results start in PENDING QA review
                    evidence_id=batch.evidence_id,
                )
                db.add(analysis)
                await db.flush()
                analyses_created += 1

                for r, prev_res in rows_to_create_with_prev:
                    r_mapped = r.mapped_payload or {}
                    norm_meth = r_mapped.get("normalization_method") or "DETERMINISTIC_SCALING"
                    norm_ver = r_mapped.get("normalization_version") or "UNIT_CONV_V1.0"

                    new_result = LaboratoryResult(
                        id=uuid.uuid4(),
                        analysis_id=analysis.id,
                        physical_sample_id=s_id,
                        analyte=r.canonical_analyte,
                        raw_value=r.raw_value,
                        raw_unit=r.raw_unit,
                        normalized_value=r.normalized_value,
                        normalized_unit=r.normalized_unit,
                        normalization_method=norm_meth,
                        normalization_version=norm_ver,
                        is_superseded=False,
                        supersedes_id=prev_res.id if prev_res else None,
                    )
                    db.add(new_result)
                    await db.flush()

                    if prev_res:
                        prev_res.is_superseded = True
                        prev_res.superseded_by_id = new_result.id

                    r.resulting_lab_result_id = new_result.id
                    results_created += 1

            sample_obj.status = "ANALYZED"
            sample_obj.updated_at = now_utc
            samples_analyzed.add(sample_obj.id)

        # Mark skipped rows
        for sr in skipped_rows:
            sr.validation_status = "SKIPPED"

        batch.imported_rows = len(eligible_rows) - exact_duplicates_skipped
        batch.skipped_rows = len(skipped_rows) + exact_duplicates_skipped
        batch.status = "IMPORTED" if len(skipped_rows) == 0 and exact_duplicates_skipped == 0 else "PARTIALLY_IMPORTED"
        batch.updated_at = now_utc

        await db.flush()

        return {
            "batch_id": batch.id,
            "status": batch.status,
            "imported_rows": batch.imported_rows,
            "skipped_rows": batch.skipped_rows,
            "analyses_created": analyses_created,
            "results_created": results_created,
            "samples_analyzed": len(samples_analyzed),
            "message": f"Successfully committed batch: {results_created} results created across {analyses_created} analyses ({exact_duplicates_skipped} exact duplicates skipped).",
        }

    @classmethod
    async def export_errors_csv(
        cls,
        db: AsyncSession,
        batch_id: uuid.UUID,
        organization_id: uuid.UUID,
    ) -> Tuple[bytes, str]:
        """
        Exports all error and warning rows in CSV format with formula injection defenses.
        """
        batch_stmt = select(LaboratoryImportBatch).where(
            and_(
                LaboratoryImportBatch.id == batch_id,
                LaboratoryImportBatch.organization_id == organization_id,
            )
        )
        batch = (await db.execute(batch_stmt)).scalars().first()
        if not batch:
            raise HTTPException(status_code=404, detail="Batch not found.")

        rows_stmt = (
            select(LaboratoryImportRow)
            .where(
                and_(
                    LaboratoryImportRow.import_batch_id == batch.id,
                    LaboratoryImportRow.validation_status.in_(["ERROR", "WARNING", "SKIPPED"]),
                )
            )
            .order_by(LaboratoryImportRow.source_row_number.asc())
        )
        rows = (await db.execute(rows_stmt)).scalars().all()

        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

        headers = [
            "row_number",
            "validation_status",
            "validation_errors",
            "sample_code",
            "analyte",
            "raw_value",
            "raw_unit",
            "original_payload",
        ]
        writer.writerow(headers)

        for r in rows:
            msgs = "; ".join([m.get("message", "") for m in (r.validation_messages or [])])
            p = r.raw_row_payload or {}
            writer.writerow([
                r.source_row_number,
                r.validation_status,
                cls.sanitize_formula_injection(msgs),
                cls.sanitize_formula_injection(r.matched_sample_code or p.get("sample_code", "")),
                cls.sanitize_formula_injection(r.canonical_analyte or p.get("analyte", "")),
                cls.sanitize_formula_injection(r.raw_value or p.get("raw_value", "")),
                cls.sanitize_formula_injection(r.raw_unit or p.get("raw_unit", "")),
                cls.sanitize_formula_injection(str(p)),
            ])

        content = output.getvalue().encode("utf-8")
        filename = f"lab_import_errors_batch_{batch.id}.csv"
        return content, filename
